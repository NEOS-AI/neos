"""
PowerPoint 파일 파싱 유틸리티

python-pptx를 사용하여 PowerPoint 파일에서 텍스트, 이미지, 노트를 추출합니다.
"""

from typing import Dict, Any, List, Optional
import io


class PPTParser:
    """
    PowerPoint 파서

    python-pptx를 사용하여 PPTX 파일을 파싱하고 슬라이드 내용을 추출합니다.
    """

    def __init__(self):
        self.supports_pptx = self._check_pptx()

    def _check_pptx(self) -> bool:
        """python-pptx 설치 여부 확인"""
        try:
            import pptx  # noqa: F401
            return True
        except ImportError:
            return False

    def is_available(self) -> bool:
        """파서 사용 가능 여부"""
        return self.supports_pptx

    async def parse(
        self,
        file_content: Optional[bytes] = None,
        file_path: Optional[str] = None,
        max_slides: int = 1000
    ) -> Dict[str, Any]:
        """
        PowerPoint 파일 파싱

        Args:
            file_content: PPT 파일 바이너리 내용
            file_path: PPT 파일 경로
            max_slides: 최대 읽을 슬라이드 수

        Returns:
            파싱 결과 딕셔너리

        Raises:
            ImportError: python-pptx가 설치되지 않은 경우
            ValueError: 파일 내용이나 경로가 제공되지 않은 경우
        """
        if not self.is_available():
            raise ImportError(
                "python-pptx is not installed. Install it with: pip install python-pptx"
            )

        if not file_content and not file_path:
            raise ValueError("Either file_content or file_path must be provided")

        try:
            from pptx import Presentation

            # PowerPoint 파일 열기
            if file_content:
                presentation = Presentation(io.BytesIO(file_content))
            else:
                presentation = Presentation(file_path)

            # 슬라이드 정보 추출
            slides_data = self._extract_slides(presentation, max_slides)

            # 텍스트 생성
            text = self._generate_text_summary(slides_data)

            # 메타데이터 추출
            metadata = self._extract_metadata(presentation)

            result = {
                "text": text,
                "slide_count": len(presentation.slides),
                "slides": slides_data,
                "has_images": any(s.get("image_count", 0) > 0 for s in slides_data),
                "has_tables": any(s.get("table_count", 0) > 0 for s in slides_data),
                "has_notes": any(s.get("has_notes", False) for s in slides_data),
                "metadata": metadata,
                "total_shapes": sum(s.get("shape_count", 0) for s in slides_data),
                "extraction_method": "python-pptx",
            }

            return result

        except Exception as e:
            return {
                "text": "",
                "slide_count": 0,
                "error": str(e),
                "extraction_method": "python-pptx",
            }

    def _extract_slides(self, presentation, max_slides: int) -> List[Dict[str, Any]]:
        """슬라이드별 데이터 추출"""
        slides_data = []

        try:
            slides_to_read = min(len(presentation.slides), max_slides)

            for slide_idx, slide in enumerate(presentation.slides):
                if slide_idx >= slides_to_read:
                    break
                try:
                    slide_info = {
                        "slide_number": slide_idx + 1,
                        "title": "",
                        "content": [],
                        "notes": "",
                        "shape_count": 0,
                        "image_count": 0,
                        "table_count": 0,
                        "has_notes": False,
                    }

                    # 슬라이드 shape 개수
                    try:
                        slide_info["shape_count"] = len(slide.shapes)
                    except Exception:
                        pass

                    # 슬라이드 제목 추출
                    try:
                        if hasattr(slide.shapes, "title") and slide.shapes.title:
                            slide_info["title"] = slide.shapes.title.text
                    except Exception:
                        pass

                    # 슬라이드 내용 추출
                    try:
                        for shape in slide.shapes:
                            try:
                                # 텍스트가 있는 shape
                                if hasattr(shape, "text") and shape.text:
                                    text = shape.text.strip()
                                    if text and text != slide_info["title"]:
                                        slide_info["content"].append(text)

                                # 표 감지
                                if hasattr(shape, "shape_type") and shape.shape_type == 19:  # MSO_SHAPE_TYPE.TABLE
                                    slide_info["table_count"] += 1
                                    # 표 내용 추출
                                    if hasattr(shape, "table"):
                                        table_text = self._extract_table_text(shape.table)
                                        if table_text:
                                            slide_info["content"].append(f"[Table]\n{table_text}")

                                # 이미지 감지
                                if hasattr(shape, "shape_type") and shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                                    slide_info["image_count"] += 1
                            except Exception:
                                # 개별 shape 처리 실패는 무시
                                continue
                    except Exception:
                        pass

                    # 슬라이드 노트 추출 - 노트는 선택사항이므로 에러 무시
                    # (notes_slide 접근 시 내부적으로 복잡한 처리가 있어서 에러 발생 가능)
                    try:
                        if slide.has_notes_slide:
                            notes_slide = slide.notes_slide
                            text_frame = notes_slide.notes_text_frame
                            if text_frame and text_frame.text:
                                notes_text = text_frame.text.strip()
                                if notes_text:
                                    slide_info["notes"] = notes_text
                                    slide_info["has_notes"] = True
                    except:
                        # 노트 추출 실패는 무시
                        pass

                    slides_data.append(slide_info)

                except Exception as e:
                    # 개별 슬라이드 처리 실패
                    slides_data.append({
                        "slide_number": slide_idx + 1,
                        "error": str(e)
                    })

        except Exception as e:
            slides_data.append({"error": f"Slides extraction failed: {str(e)}"})

        return slides_data

    def _extract_table_text(self, table) -> str:
        """표에서 텍스트 추출"""
        try:
            rows = []
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    cells.append(cell.text.strip())
                rows.append(" | ".join(cells))
            return "\n".join(rows)
        except Exception:
            return ""

    def _generate_text_summary(self, slides_data: List[Dict[str, Any]]) -> str:
        """슬라이드 데이터에서 텍스트 요약 생성"""
        text_parts = []

        for slide in slides_data:
            if "error" in slide:
                continue

            slide_num = slide.get("slide_number", "?")
            text_parts.append(f"=== Slide {slide_num} ===")

            # 제목
            title = slide.get("title", "")
            if title:
                text_parts.append(f"Title: {title}")

            # 내용
            content = slide.get("content", [])
            if content:
                for item in content:
                    text_parts.append(item)

            # 노트
            notes = slide.get("notes", "")
            if notes:
                text_parts.append(f"\nNotes: {notes}")

            # 이미지/표 정보
            image_count = slide.get("image_count", 0)
            table_count = slide.get("table_count", 0)
            if image_count > 0:
                text_parts.append(f"[Contains {image_count} image(s)]")
            if table_count > 0:
                text_parts.append(f"[Contains {table_count} table(s)]")

            text_parts.append("")  # 슬라이드 간 구분

        return "\n".join(text_parts)

    def _extract_metadata(self, presentation) -> Dict[str, Any]:
        """PowerPoint 메타데이터 추출"""
        metadata = {}

        try:
            core_props = presentation.core_properties

            if core_props:
                metadata_fields = {
                    'title': core_props.title,
                    'author': core_props.author,
                    'subject': core_props.subject,
                    'keywords': core_props.keywords,
                    'comments': core_props.comments,
                    'category': core_props.category,
                    'created': core_props.created,
                    'modified': core_props.modified,
                    'last_modified_by': core_props.last_modified_by,
                    'revision': core_props.revision,
                }

                for key, value in metadata_fields.items():
                    if value:
                        # datetime 객체를 문자열로 변환
                        if hasattr(value, 'isoformat'):
                            value = value.isoformat()
                        metadata[key] = value

        except Exception as e:
            metadata['extraction_error'] = str(e)

        return metadata


# 편의 함수
async def parse_ppt(
    file_content: Optional[bytes] = None,
    file_path: Optional[str] = None,
    max_slides: int = 1000
) -> Dict[str, Any]:
    """
    PowerPoint 파일 파싱 편의 함수

    Args:
        file_content: PPT 바이너리 내용
        file_path: PPT 파일 경로
        max_slides: 최대 슬라이드 수

    Returns:
        파싱 결과
    """
    parser = PPTParser()
    return await parser.parse(
        file_content=file_content,
        file_path=file_path,
        max_slides=max_slides
    )
