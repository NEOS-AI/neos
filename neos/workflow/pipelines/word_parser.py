"""
Word 문서 파싱 유틸리티

python-docx를 사용하여 Word 문서에서 텍스트, 테이블, 스타일 정보를 추출합니다.
"""

from typing import Dict, Any, List, Optional
import io
from pathlib import Path


class WordParser:
    """
    Word 문서 파서

    python-docx를 사용하여 DOCX 파일을 파싱하고 텍스트와 구조를 추출합니다.
    """

    def __init__(self):
        self.supports_docx = self._check_docx()

    def _check_docx(self) -> bool:
        """python-docx 설치 여부 확인"""
        try:
            import docx
            return True
        except ImportError:
            return False

    def is_available(self) -> bool:
        """파서 사용 가능 여부"""
        return self.supports_docx

    async def parse(
        self,
        file_content: Optional[bytes] = None,
        file_path: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        Word 문서 파싱

        Args:
            file_content: Word 파일 바이너리 내용
            file_path: Word 파일 경로

        Returns:
            파싱 결과 딕셔너리

        Raises:
            ImportError: python-docx가 설치되지 않은 경우
            ValueError: 파일 내용이나 경로가 제공되지 않은 경우
        """
        if not self.is_available():
            raise ImportError(
                "python-docx is not installed. Install it with: pip install python-docx"
            )

        if not file_content and not file_path:
            raise ValueError("Either file_content or file_path must be provided")

        try:
            from docx import Document

            # Word 문서 열기
            if file_content:
                doc = Document(io.BytesIO(file_content))
            else:
                doc = Document(file_path)

            # 텍스트 추출
            text_result = self._extract_text(doc)

            # 테이블 추출
            tables = self._extract_tables(doc)

            # 메타데이터 추출
            metadata = self._extract_metadata(doc)

            # 스타일 정보
            styles_info = self._extract_styles_info(doc)

            result = {
                "text": text_result["text"],
                "paragraph_count": text_result["paragraph_count"],
                "char_count": len(text_result["text"]),
                "word_count": len(text_result["text"].split()),
                "tables": tables,
                "table_count": len(tables),
                "has_images": text_result.get("has_images", False),
                "has_tables": len(tables) > 0,
                "metadata": metadata,
                "styles": styles_info,
                "extraction_method": "python-docx",
            }

            return result

        except Exception as e:
            return {
                "text": "",
                "paragraph_count": 0,
                "error": str(e),
                "extraction_method": "python-docx",
            }

    def _extract_text(self, doc) -> Dict[str, Any]:
        """Word 문서에서 텍스트 추출"""
        paragraphs = []
        has_images = False

        try:
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    paragraphs.append(text)

                # 이미지 감지 (간단한 휴리스틱)
                if para._element.xpath('.//pic:pic'):
                    has_images = True

            combined_text = "\n\n".join(paragraphs)

            return {
                "text": combined_text,
                "paragraph_count": len(paragraphs),
                "has_images": has_images,
            }

        except Exception as e:
            return {
                "text": f"Text extraction error: {str(e)}",
                "paragraph_count": 0,
                "has_images": False,
            }

    def _extract_tables(self, doc) -> List[Dict[str, Any]]:
        """Word 문서에서 테이블 추출"""
        tables_data = []

        try:
            for table_idx, table in enumerate(doc.tables):
                table_info = {
                    "table_number": table_idx + 1,
                    "row_count": len(table.rows),
                    "col_count": len(table.columns) if table.rows else 0,
                    "data": []
                }

                # 테이블 데이터 추출
                for row in table.rows:
                    row_data = []
                    for cell in row.cells:
                        row_data.append(cell.text.strip())
                    table_info["data"].append(row_data)

                tables_data.append(table_info)

        except Exception as e:
            tables_data.append({"error": str(e)})

        return tables_data

    def _extract_metadata(self, doc) -> Dict[str, Any]:
        """Word 문서 메타데이터 추출"""
        metadata = {}

        try:
            core_properties = doc.core_properties

            # 주요 메타데이터 필드
            metadata_fields = {
                'title': core_properties.title,
                'author': core_properties.author,
                'subject': core_properties.subject,
                'keywords': core_properties.keywords,
                'comments': core_properties.comments,
                'last_modified_by': core_properties.last_modified_by,
                'created': core_properties.created,
                'modified': core_properties.modified,
                'category': core_properties.category,
            }

            for key, value in metadata_fields.items():
                if value:
                    # datetime 객체를 문자열로 변환
                    if hasattr(value, 'isoformat'):
                        value = value.isoformat()
                    metadata[key] = value

        except Exception as e:
            metadata['extraction_error'] = str(e)

        return metadata

    def _extract_styles_info(self, doc) -> Dict[str, Any]:
        """Word 문서 스타일 정보 추출"""
        styles_info = {
            "headings": [],
            "has_bold": False,
            "has_italic": False,
        }

        try:
            for para in doc.paragraphs:
                # 제목 스타일 감지
                if para.style.name.startswith('Heading'):
                    styles_info["headings"].append({
                        "level": para.style.name,
                        "text": para.text[:100]  # 처음 100자만
                    })

                # 볼드/이탤릭 감지
                for run in para.runs:
                    if run.bold:
                        styles_info["has_bold"] = True
                    if run.italic:
                        styles_info["has_italic"] = True

        except Exception as e:
            styles_info["error"] = str(e)

        return styles_info


# 편의 함수
async def parse_word(
    file_content: Optional[bytes] = None,
    file_path: Optional[str] = None
) -> Dict[str, Any]:
    """
    Word 문서 파싱 편의 함수

    Args:
        file_content: Word 바이너리 내용
        file_path: Word 파일 경로

    Returns:
        파싱 결과
    """
    parser = WordParser()
    return await parser.parse(file_content=file_content, file_path=file_path)
